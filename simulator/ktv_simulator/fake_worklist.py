"""TẠM: sinh message input giả đúng file ``API-Goi-y-cong-viec.xlsx`` và file Excel để review.

Chưa có dữ liệu thật theo API nên ghép từ hai nguồn:
- THẬT: việc bảo trì, KTV, tọa độ check-in, nhịp sự kiện của luồng sự kiện HNI_04.
- GIẢ: 4 nhóm còn lại (triển khai, thu hồi, hóa đơn, onsite), giờ hẹn, ca làm... theo
  các giả định trong ``ASSUMPTIONS``. Mọi giả định được liệt kê ở sheet "01 · Giả định".

Mỗi message là input của một lần gọi AI cho một KTV, cộng 3 field vỏ Kafka đề xuất
(``message_id``, ``planned_at``, ``trigger``). Message sinh tại: 06:00 đầu ngày, lúc KTV xong
việc (checkout/đóng), lúc có việc mới, lúc khách đổi giờ.

    PYTHONPATH=src:simulator .venv/bin/python -m ktv_simulator.fake_worklist \\
        --events data/sample/events_2026-06_HNI_04.jsonl \\
        --checkins data/sample/QOS_MAINT_CHECKIN_INFO_utf8.csv \\
        --day 2026-06-08 --day 2026-06-16 --day 2026-06-21 --out-dir artifacts/fake
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
import re
import sys
import zlib
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from statistics import median

from ktv_routing.contract import GeoPoint, JobFilter, JobState, WorkloadQuery
from ktv_routing.travel import distance_km

from .provider import EventWorkloadProvider

GROUP_KEYS = ("trien_khai", "bao_tri", "thu_hoi", "hoa_don", "onsite")  # Thứ tự khóa của tasks.
GROUP_IDS = {name: index for index, name in enumerate(GROUP_KEYS, start=1)}
CHECKIN_BEFORE_END = "Check-in trước mốc hẹn cuối"


@dataclass(frozen=True, slots=True)
class TaskKind:
    group: str
    type_id: int
    name: str
    sla_minutes: int | None
    on_time: str
    priority: int
    handle: tuple[int, int]  # Thời gian xử lý giả: (trung bình, độ lệch) phút.
    appointment_rate: float  # Tỷ lệ việc có giờ hẹn.
    weight: int  # Tỷ trọng trong nhóm.


# Sheet 05 file API. type_id ngoài mẫu (box_cam_only, swap, giao_thiet_bi_cam, bao_tri_logic,
# chap_chon_suy_hao) là số tạm, chờ chốt bảng số. Handle, tỷ lệ hẹn, tỷ trọng là GIẢ.
KINDS = (
    TaskKind("trien_khai", 3, "trien_khai_net", 120, CHECKIN_BEFORE_END, 3, (120, 20), 1.0, 50),
    TaskKind("trien_khai", 4, "box_cam_only", 120, CHECKIN_BEFORE_END, 3, (60, 15), 1.0, 20),
    TaskKind("trien_khai", 5, "swap", 60, CHECKIN_BEFORE_END, 3, (45, 10), 1.0, 20),
    TaskKind("trien_khai", 6, "giao_thiet_bi_cam", 60, "Hoàn tất trong ngày hẹn", 3, (30, 10), 1.0, 10),
    TaskKind("bao_tri", 1, "bao_tri_vat_ly", 60, CHECKIN_BEFORE_END, 1, (60, 15), 1.0, 60),
    TaskKind("bao_tri", 2, "bao_tri_logic", 60, CHECKIN_BEFORE_END, 2, (45, 15), 1.0, 40),
    TaskKind("thu_hoi", 1, "thu_hoi_thiet_bi", None, "Hoàn tất trong tháng", 4, (20, 5), 0.1, 100),
    TaskKind("hoa_don", 1, "hoa_don_tra_truoc", None, "Hoàn tất trong tháng", 4, (15, 5), 0.05, 40),
    TaskKind("hoa_don", 2, "hoa_don_tra_sau", None, "Hoàn tất trong tháng", 4, (15, 5), 0.05, 60),
    TaskKind("onsite", 1, "phieu_onsite", None, "Hoàn tất trong tháng", 2, (45, 15), 0.4, 40),
    TaskKind("onsite", 2, "ngung_ket_noi_4h", None, "Hoàn tất trong ngày tạo phiếu", 2, (30, 10), 0.0, 30),
    TaskKind("onsite", 3, "chap_chon_suy_hao", None, "Hoàn tất trong ngày tạo phiếu", 4, (40, 10), 0.2, 30),
)
KIND_BY_NAME = {(kind.group, kind.name): kind for kind in KINDS}
GSAFE = (12, "gsafe")  # Loại con, như JSON mẫu: trien_khai_net + gsafe.
PENDING_STATUS = (6, "check_in")  # Trạng thái trong JSON mẫu.
IN_PROGRESS_STATUS = 10

# Mọi số bịa nằm ở đây; sheet "01 · Giả định" in lại bảng này.
SHIFT = "08:00-17:30"
OVERTIME = "17:30-21:00"
OT_RATE = 0.2
EXTRA_TASKS_PER_DAY = 4.0  # Việc 4 nhóm giả / KTV / ngày (trung bình, Poisson).
GROUP_WEIGHTS = {"trien_khai": 25, "thu_hoi": 10, "hoa_don": 20, "onsite": 10}
GSAFE_RATE = 0.1
VAT_LY_RATE = 0.6
CARRY_OVER_RATE = 0.5  # Việc giả tạo từ chiều hôm trước.
MAIN_PLOT_RATE = 0.85  # Việc giả nằm ở lô chính (dữ liệu thật: trung vị 85%).
SUPPORT_MIN_SHARE = 0.03
MAX_SUPPORT_PLOTS = 2
JITTER_M = 100
MISSING_LATLNG_RATE = 0.01
OVERLAP_RATE = 0.1
RESCHEDULE_RATE = 0.1
BLOCK_ID = 4
FIRST_PLOT_ID = 101
FAKE_TASK_ID = 9_000_000
FAKE_STAFF_ID = 90_000_000
DAY_START = time(6, 0)
DAY_END = time(21, 0)
TRIGGERS = ("DAY_START", "TASK_DONE", "TASK_NEW", "RESCHEDULE")

# (Chủ đề, field trong API, cách giả lập — viết cho người không đọc code, ví dụ, nguồn).
# Nguồn: "Thật" = lấy nguyên dữ liệu thật; "Suy ra" = tính từ dữ liệu thật; "Tự đặt" = con số nhóm tự chọn, cần duyệt.
ASSUMPTIONS = (
    ("Việc bảo trì", "tasks.bao_tri",
     "Dùng đúng các việc bảo trì còn mở của KTV tại thời điểm đó, lấy từ dữ liệu thật tháng 6/2026 chi nhánh HNI_04.",
     "KTV TIN0401.VINHNT10 lúc 9h còn 5 việc bảo trì chưa làm", "Thật"),
    ("Việc bảo trì: vật lý hay logic", "task_type",
     "Dữ liệu cũ không ghi rõ bảo trì vật lý hay logic, nên chia ngẫu nhiên nhưng cố định cho từng việc.",
     f"{VAT_LY_RATE:.0%} vật lý (ưu tiên 1), {1 - VAT_LY_RATE:.0%} logic (ưu tiên 2)", "Tự đặt"),
    ("Việc bảo trì: giờ hẹn khách", "appointment",
     "Dữ liệu cũ không có giờ hẹn. Đặt giờ hẹn = giờ tạo phiếu + 23 tiếng, nên hạn check-in (hẹn + 60 phút) rơi vào khoảng 24 tiếng sau khi tạo phiếu, "
     "khớp với cờ đúng hẹn trong dữ liệu cũ ở 90% trường hợp. Hẹn rơi ngoài giờ làm thì dời sang 8h–10h sáng.",
     "Tạo phiếu 9h10 ngày 15 → hẹn 8h30 ngày 16, phải check-in trước 9h30", "Tự đặt"),
    ("Việc bảo trì: vị trí khách", "latlng, location",
     "Địa chỉ trong phiếu; tọa độ lấy từ lần check-in đầu tiên của KTV tại nhà khách, tức vị trí khách thật. Việc chưa ai tới thì dùng tâm phường.",
     "—", "Thật"),
    ("Việc bảo trì: thời gian xử lý", "handle_minutes",
     "Để trống, theo đúng file API: hệ thống AI tự dùng định mức theo loại việc.",
     "—", "Suy ra"),
    ("4 nhóm việc chưa có dữ liệu", "tasks.trien_khai, thu_hoi, hoa_don, onsite",
     f"Chưa có dữ liệu thật của triển khai, thu hồi, hóa đơn, onsite, nên mỗi KTV được tạo thêm trung bình {EXTRA_TASKS_PER_DAY:g} việc/ngày "
     "(ngày nhiều ngày ít, ngẫu nhiên).",
     f"Tỷ lệ giữa 4 nhóm: triển khai {GROUP_WEIGHTS['trien_khai']} · hóa đơn {GROUP_WEIGHTS['hoa_don']} · "
     f"thu hồi {GROUP_WEIGHTS['thu_hoi']} · onsite {GROUP_WEIGHTS['onsite']}", "Tự đặt"),
    ("4 nhóm tạo thêm: loại việc", "task_type, task_sub",
     "Trong mỗi nhóm chia theo tỷ lệ bên cạnh. SLA và mức ưu tiên lấy đúng bảng SLA của file API.",
     "Triển khai: mới NET 50 · box/cam 20 · swap 20 · giao cam 10; hóa đơn: trả sau 60 · trả trước 40; "
     f"onsite: phiếu onsite 40 · ngưng kết nối 30 · chập chờn 30; {GSAFE_RATE:.0%} việc triển khai mới có gói Gsafe", "Tự đặt"),
    ("4 nhóm tạo thêm: vị trí khách", "latlng, location",
     f"Lấy ngẫu nhiên một địa chỉ khách thật trong địa bàn của KTV rồi lệch đi dưới {JITTER_M} m. "
     f"{MAIN_PLOT_RATE:.0%} việc nằm ở địa bàn chính, còn lại ở địa bàn hỗ trợ.",
     "—", "Tự đặt"),
    ("4 nhóm tạo thêm: thời gian xử lý", "handle_minutes",
     "Dao động quanh mức trung bình của từng loại việc.",
     "Triển khai mới ~120 phút · box/cam ~60 · swap ~45 · giao cam ~30 · thu hồi ~20 · hóa đơn ~15 · onsite 30–45", "Tự đặt"),
    ("4 nhóm tạo thêm: giờ hẹn", "appointment",
     "Việc triển khai luôn có giờ hẹn; thu hồi, hóa đơn gần như không hẹn; onsite hẹn một phần. Giờ hẹn nằm trong 8h–16h30, "
     "ít nhất 1 tiếng sau khi tạo phiếu, làm tròn 30 phút.",
     "Onsite: phiếu onsite 40% có hẹn, chập chờn 20%, ngưng kết nối 0%", "Tự đặt"),
    ("4 nhóm tạo thêm: lúc phát sinh và lúc xong", "—",
     f"{CARRY_OVER_RATE:.0%} việc có từ chiều hôm trước, còn lại phát sinh trong ngày (7h–15h). Mỗi việc được coi là làm xong "
     "vào một thời điểm hợp lý trong ca; không kịp trước 17h30 thì để sang hôm sau.",
     "—", "Tự đặt"),
    ("Tình huống khó được cố ý đưa vào", "—",
     f"Để kiểm tra AI xử lý tình huống xấu: {MISSING_LATLNG_RATE:.0%} việc thiếu tọa độ, {OVERLAP_RATE:.0%} việc có hẹn trùng giờ với việc khác, "
     f"{RESCHEDULE_RATE:.0%} việc có hẹn bị khách dời thêm 1–2 tiếng trong ngày.",
     "—", "Tự đặt"),
    ("Địa bàn (lô)", "staff.plots, task_plots_id",
     "Mỗi phường là một địa bàn. Địa bàn chính của KTV là phường KTV làm nhiều việc nhất trong tháng; "
     "địa bàn hỗ trợ là tối đa 2 phường tiếp theo mà KTV thường sang giúp.",
     "Dữ liệu thật: trung bình 85% việc của một KTV nằm trong địa bàn chính", "Suy ra"),
    ("Việc thuộc địa bàn nào của KTV", "staff_plots_id, staff_role",
     "Việc nằm trong địa bàn của KTV thì ghi địa bàn đó (chính hoặc hỗ trợ). Việc ở ngoài thì ghi địa bàn gần nhất của KTV.",
     "—", "Suy ra"),
    ("Khối / vùng", "block_id",
     "Chưa có danh mục khối, tạm coi cả chi nhánh là một khối.", str(BLOCK_ID), "Tự đặt"),
    ("location_id", "location_id",
     "Field này có trong JSON mẫu nhưng chưa có mô tả, tạm ghi bằng mã địa bàn của việc.", "—", "Tự đặt"),
    ("Mã loại việc chưa có số", "task_type_id",
     "File API chỉ cho số của vài loại, các loại còn lại được đánh số tạm, chờ chốt bảng số.",
     "box/cam 4 · swap 5 · giao cam 6 · bảo trì logic 2 · chập chờn 3", "Tự đặt"),
    ("Trạng thái việc", "task_status",
     "Việc chờ làm dùng trạng thái như JSON mẫu; việc KTV đang làm dở đặt ở current_task với trạng thái 10.",
     f"{PENDING_STATUS[0]} = {PENDING_STATUS[1]} · {IN_PROGRESS_STATUS} = đang thực hiện", "Tự đặt"),
    ("Mã nhân viên", "staff_id",
     "Mã nhân viên thật của KTV trong dữ liệu check-in; KTV không có thì tạo mã.", "00324668", "Thật"),
    ("Vị trí hiện tại của KTV", "staff.latlng",
     "Chi nhánh này không có GPS, nên dùng vị trí check-in hoặc checkout gần nhất của KTV. Chưa có thì lấy tâm địa bàn chính.",
     "—", "Suy ra"),
    ("Việc KTV đang làm dở", "current_task",
     "Việc bảo trì KTV đã check-in mà chưa checkout tại thời điểm đó.", "—", "Thật"),
    ("Ca làm", "available",
     f"Mọi KTV làm {SHIFT.replace('-', '–')}; {OT_RATE:.0%} KTV có đăng ký tăng ca {OVERTIME.replace('-', '–')}.",
     f"{SHIFT},{OVERTIME}", "Tự đặt"),
    ("Khi nào gửi dữ liệu cho AI", "message_id, planned_at, trigger",
     "Mỗi bản ghi là một lần gửi danh sách việc của một KTV cho AI. Gửi lúc 6h sáng, mỗi khi KTV làm xong việc, có việc mới, "
     "hoặc khách đổi hẹn. 3 field này là đề xuất thêm để đóng gói khi gửi qua Kafka, không có trong file API.",
     "DAY_START · TASK_DONE · TASK_NEW · RESCHEDULE", "Tự đặt"),
)


def _dt(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _latlng(point: GeoPoint | None) -> str:
    return "" if point is None else f"{point.lat:.6f},{point.lng:.6f}"


def _ceil30(value: datetime) -> datetime:
    value = value.replace(second=0, microsecond=0)
    extra = (-value.minute) % 30
    return value + timedelta(minutes=extra)


def _stable(text: str) -> float:
    """Số trong [0, 1) cố định theo chuỗi, để việc thật luôn được gán cùng giá trị."""

    return zlib.crc32(text.encode()) / 2**32


def _poisson(rng: random.Random, mean: float) -> int:
    limit, count, product = math.exp(-mean), 0, rng.random()
    while product > limit:
        count += 1
        product *= rng.random()
    return count


def _weighted(rng: random.Random, items: dict):
    return rng.choices(list(items), weights=list(items.values()))[0]


@dataclass(frozen=True, slots=True)
class Plot:
    id: int
    name: str
    center: GeoPoint


@dataclass(slots=True)
class Context:
    """Những gì lấy được từ dữ liệu thật, dùng chung cho mọi message."""

    plots: dict[str, Plot]  # Phường → lô.
    staff_plots: dict[str, list[tuple[Plot, int, float]]]  # KTV → [(lô, role, tỷ lệ việc)].
    staff_ids: dict[str, tuple[str, str]]  # KTV → (staff_id, nguồn).
    points: dict[str, GeoPoint]  # Việc → tọa độ check-in đầu tiên.
    created: dict[str, datetime]
    area: dict[str, str]  # Việc → phường.
    customers: dict[str, list[tuple[str, GeoPoint]]]  # Phường → [(địa chỉ, tọa độ)] khách thật.
    numeric_ids: dict[str, int] = field(default_factory=dict)

    def task_id(self, job_id: str) -> int:
        if job_id.isdigit():
            return int(job_id)
        return self.numeric_ids.setdefault(job_id, 8_000_000 + len(self.numeric_ids))


def load_context(events_path: Path, checkins_csv: Path, encoding: str = "utf-8-sig") -> Context:
    owner: dict[str, str] = {}
    created: dict[str, datetime] = {}
    job_area: dict[str, str] = {}
    job_point: dict[str, GeoPoint] = {}
    address: dict[str, str] = {}
    with events_path.open(encoding="utf-8") as handle:
        next(handle)
        for line in handle:
            if '"JOB_CREATED"' not in line:
                continue
            event = json.loads(line)
            job_id = event["job_id"]
            created[job_id] = datetime.fromisoformat(event["at"])
            if event.get("emp_account"):
                owner[job_id] = event["emp_account"]
            if event.get("area"):
                job_area[job_id] = event["area"]
            if event.get("location"):
                job_point[job_id] = GeoPoint(event["location"]["lat"], event["location"]["lng"])
            if event.get("address"):
                address[job_id] = event["address"]

    csv.field_size_limit(sys.maxsize)
    points: dict[str, GeoPoint] = {}
    codes: dict[str, Counter] = defaultdict(Counter)
    pattern = re.compile(r"(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)")
    with checkins_csv.open(encoding=encoding, newline="") as handle:
        for row in csv.DictReader(handle):
            job_id = (row.get("CHECKLIST_ID") or "").strip()
            match = pattern.search(row.get("LAT_LNG_IN") or "")
            if match and job_id not in points:
                lat, lng = float(match[1]), float(match[2])
                if 8 <= lat <= 24 and 102 <= lng <= 110:
                    points[job_id] = GeoPoint(lat, lng)
            code = (row.get("EMP_CODE") or "").strip()
            if code and job_id in owner:
                codes[owner[job_id]][code] += 1

    by_area: dict[str, list[GeoPoint]] = defaultdict(list)
    customers: dict[str, list[tuple[str, GeoPoint]]] = defaultdict(list)
    for job_id, area in job_area.items():
        point = points.get(job_id) or job_point.get(job_id)
        if point is not None:
            by_area[area].append(point)
        if job_id in points:
            customers[area].append((address.get(job_id, ""), points[job_id]))
    plots = {
        area: Plot(
            FIRST_PLOT_ID + index,
            area,
            GeoPoint(median(p.lat for p in by_area[area]), median(p.lng for p in by_area[area])),
        )
        for index, area in enumerate(sorted(by_area))
    }

    counts: dict[str, Counter] = defaultdict(Counter)
    for job_id, account in owner.items():
        if job_area.get(job_id) in plots:
            counts[account][job_area[job_id]] += 1
    staff_plots = {}
    for account, counter in counts.items():
        total = sum(counter.values())
        ranked = counter.most_common()
        chosen = [(plots[ranked[0][0]], 1, ranked[0][1] / total)]
        chosen += [
            (plots[area], 2, count / total)
            for area, count in ranked[1 : 1 + MAX_SUPPORT_PLOTS]
            if count / total >= SUPPORT_MIN_SHARE
        ]
        staff_plots[account] = chosen

    staff_ids = {}
    for index, account in enumerate(sorted(set(owner.values()))):
        if codes.get(account):
            staff_ids[account] = (codes[account].most_common(1)[0][0], "THẬT")
        else:
            staff_ids[account] = (f"{FAKE_STAFF_ID + index:08d}", "GIẢ")

    return Context(plots, staff_plots, staff_ids, points, created, job_area, dict(customers))


# ---------------------------------------------------------------- việc giả


@dataclass(slots=True)
class FakeTask:
    task_id: int
    kind: TaskKind
    sub: tuple[int, str]
    created: datetime
    appointment: datetime | None
    handle: int
    plot: Plot
    address: str
    point: GeoPoint | None
    done: datetime | None = None  # None: chưa xong trong ngày.
    reschedule: tuple[datetime, datetime] | None = None  # (lúc khách báo, giờ hẹn mới).

    def appointment_at(self, at: datetime) -> datetime | None:
        if self.reschedule is not None and at >= self.reschedule[0]:
            return self.reschedule[1]
        return self.appointment


def _schedule_done(task: FakeTask, day: date, rng: random.Random) -> None:
    opens = datetime.combine(day, time(8, 0))
    start = max(task.created, task.appointment or opens, opens) + timedelta(minutes=rng.uniform(0, 180))
    done = start + timedelta(minutes=task.handle)
    task.done = done.replace(second=0, microsecond=0) if done <= datetime.combine(day, time(17, 30)) else None


def fake_tasks(ctx: Context, account: str, day: date, seed: int, next_id: list[int]) -> list[FakeTask]:
    """Việc 4 nhóm giả của một KTV trong một ngày; cố định theo (seed, ngày, KTV)."""

    plots = ctx.staff_plots.get(account)
    if not plots:
        return []
    rng = random.Random(f"{seed}:{day}:{account}")
    tasks: list[FakeTask] = []
    for _ in range(_poisson(rng, EXTRA_TASKS_PER_DAY)):
        group = _weighted(rng, GROUP_WEIGHTS)
        kind = _weighted(rng, {kind: kind.weight for kind in KINDS if kind.group == group})
        sub = GSAFE if kind.name == "trien_khai_net" and rng.random() < GSAFE_RATE else (0, "")
        if rng.random() < CARRY_OVER_RATE:
            created = datetime.combine(day - timedelta(days=1), time(14)) + timedelta(minutes=rng.randrange(7 * 60))
        else:
            created = datetime.combine(day, time(7)) + timedelta(minutes=rng.randrange(8 * 60))
        appointment = None
        if rng.random() < kind.appointment_rate:
            earliest = max(created + timedelta(minutes=60), datetime.combine(day, time(8)))
            latest = datetime.combine(day, time(16, 30))
            if earliest <= latest:
                appointment = _ceil30(earliest + (latest - earliest) * rng.random())
        plot = plots[0][0] if rng.random() < MAIN_PLOT_RATE or len(plots) == 1 else rng.choice(plots[1:])[0]
        customers = ctx.customers.get(plot.name) or [("", plot.center)]
        address, point = rng.choice(customers)
        meters = rng.uniform(0, JITTER_M)
        angle = rng.uniform(0, 2 * math.pi)
        point = GeoPoint(
            point.lat + meters * math.cos(angle) / 111_195,
            point.lng + meters * math.sin(angle) / (111_195 * math.cos(math.radians(point.lat))),
        )
        if rng.random() < MISSING_LATLNG_RATE:
            point = None
        handle = max(5, round(rng.gauss(*kind.handle)))
        tasks.append(FakeTask(next_id[0], kind, sub, created, appointment, handle, plot, address, point))
        next_id[0] += 1

    booked = [task for task in tasks if task.appointment is not None]
    for task in booked:
        others = [other for other in booked if other is not task]
        if others and rng.random() < OVERLAP_RATE:
            task.appointment = rng.choice(others).appointment
    for task in tasks:
        _schedule_done(task, day, rng)
        if task.appointment is not None and rng.random() < RESCHEDULE_RATE:
            told = task.appointment - timedelta(minutes=rng.randrange(60, 181))
            if told >= max(task.created, datetime.combine(day, DAY_START)):
                task.reschedule = (told, task.appointment + timedelta(minutes=rng.choice((60, 90, 120))))
    return tasks


# ---------------------------------------------------------------- message


def _bao_tri_kind(job_id: str) -> TaskKind:
    name = "bao_tri_vat_ly" if _stable(job_id) < VAT_LY_RATE else "bao_tri_logic"
    return KIND_BY_NAME[("bao_tri", name)]


def _bao_tri_appointment(job_id: str, created: datetime) -> datetime:
    """Giờ tạo + 23 giờ; rơi ngoài 08:00–17:00 thì dời vào 08:00–10:00 (rải theo job, tránh dồn 08:00)."""

    value = _ceil30(created + timedelta(hours=23))
    morning = timedelta(minutes=30 * int(_stable(f"appt:{job_id}") * 5))
    if value.time() > time(17):
        value = datetime.combine(value.date() + timedelta(days=1), time(8)) + morning
    elif value.time() < time(8):
        value = datetime.combine(value.date(), time(8)) + morning
    return value


def _staff_plot(ctx: Context, account: str, area: str | None, point: GeoPoint | None) -> tuple[int | None, int, int]:
    """(task_plots_id, staff_plots_id, staff_role) của một việc."""

    own = ctx.staff_plots.get(account, [])
    plot = ctx.plots.get(area) if area else None
    for candidate, role, _ in own:
        if plot is not None and candidate.id == plot.id:
            return plot.id, candidate.id, role
    if not own:
        return (plot.id if plot else None), 0, 2
    anchor = plot.center if plot is not None else point
    if anchor is None:
        candidate, role, _ = own[0]
    else:
        candidate, role, _ = min(own, key=lambda item: distance_km(item[0].center, anchor))
    return (plot.id if plot else candidate.id), candidate.id, role


def _task(ctx, account, *, task_id, kind, sub, appointment, location, point, handle, area) -> dict:
    task_plot, staff_plot, role = _staff_plot(ctx, account, area, point)
    return {
        "task_id": task_id,
        "task_group_id": GROUP_IDS[kind.group],
        "task_group_name": kind.group,
        "task_type_id": kind.type_id,
        "task_type_name": kind.name,
        "task_sub_id": sub[0],
        "task_sub_name": sub[1],
        "task_status_id": PENDING_STATUS[0],
        "task_status_name": PENDING_STATUS[1],
        "sla": {"sla_minutes": kind.sla_minutes, "priority_in_day": kind.priority},
        "appointment": _dt(appointment) if appointment else "",
        "location": location or "",
        "latlng": _latlng(point),
        "handle_minutes": handle,
        "task_plots_id": task_plot,
        "staff_plots_id": staff_plot,
        "staff_role": role,
        "block_id": BLOCK_ID,
        "location_id": task_plot,
    }


def build_message(
    ctx: Context,
    message_id: str,
    planned_at: datetime,
    trigger: str,
    account: str,
    technician,
    fakes: list[FakeTask],
) -> tuple[dict, dict[int, str]]:
    """Message input của một KTV. Trả kèm nguồn từng việc: REAL, REAL_WARD (tọa độ tâm phường) hoặc FAKE."""

    sources: dict[int, str] = {}
    tasks: dict[str, list[dict]] = {key: [] for key in GROUP_KEYS}
    current = None
    jobs = list(technician.jobs) if technician is not None else []
    running = sorted(
        (job for job in jobs if job.state is JobState.IN_PROGRESS), key=lambda job: job.started_at or planned_at
    )
    if running:
        job = running[-1]
        current = {
            "task_id": ctx.task_id(job.job_id),
            "task_status_id": IN_PROGRESS_STATUS,
            "task_type_id": _bao_tri_kind(job.job_id).type_id,
        }
    for job in jobs:
        if running and job is running[-1]:
            continue
        kind = _bao_tri_kind(job.job_id)
        point = ctx.points.get(job.job_id)
        task_id = ctx.task_id(job.job_id)
        sources[task_id] = "REAL" if point is not None else "REAL_WARD"
        tasks["bao_tri"].append(
            _task(
                ctx,
                account,
                task_id=task_id,
                kind=kind,
                sub=(0, ""),
                appointment=_bao_tri_appointment(job.job_id, ctx.created.get(job.job_id, planned_at)),
                location=job.address,
                point=point or job.location,
                handle="",
                area=job.area,
            )
        )
    for fake in fakes:
        if fake.created > planned_at or (fake.done is not None and fake.done <= planned_at):
            continue
        sources[fake.task_id] = "FAKE"
        tasks[fake.kind.group].append(
            _task(
                ctx,
                account,
                task_id=fake.task_id,
                kind=fake.kind,
                sub=fake.sub,
                appointment=fake.appointment_at(planned_at),
                location=fake.address,
                point=fake.point,
                handle=fake.handle,
                area=fake.plot.name,
            )
        )
    for items in tasks.values():
        items.sort(key=lambda item: item["task_id"])

    own = ctx.staff_plots.get(account, [])
    fix = technician.last_location if technician is not None else None
    ot = _stable(f"ot:{account}") < OT_RATE
    message = {
        "message_id": message_id,
        "planned_at": _dt(planned_at),
        "trigger": trigger,
        "staff": {
            "staff_id": ctx.staff_ids.get(account, (f"{FAKE_STAFF_ID:08d}", "GIẢ"))[0],
            "staff_account": account,
            "staff_location": "HNI_04",
            "latlng": _latlng(fix.location if fix else (own[0][0].center if own else None)),
            "available": SHIFT + ("," + OVERTIME if ot else ""),
            "plots": [
                {"id": plot.id, "name": plot.name, "role": role, "block_id": BLOCK_ID} for plot, role, _ in own
            ],
            "current_task": current,
        },
        "tasks": tasks,
    }
    return message, sources


def generate(
    events_path: Path, ctx: Context, days: list[date], seed: int = 42
) -> tuple[list[dict], list[dict[int, str]]]:
    """Mọi message của các ngày, sắp theo thời gian."""

    messages: list[dict] = []
    sources: list[dict[int, str]] = []
    next_id = [FAKE_TASK_ID]
    trigger_of = {"CHECKOUT": "TASK_DONE", "JOB_CLOSED": "TASK_DONE", "JOB_CREATED": "TASK_NEW"}
    with EventWorkloadProvider(events_path) as provider:
        for day in sorted(days):
            start, end = datetime.combine(day, DAY_START), datetime.combine(day, DAY_END)
            opening = provider.build(WorkloadQuery(start))[0]
            accounts = {technician.emp_account for technician in opening.technicians}
            triggers: dict[tuple[datetime, str], str] = {}
            for change in provider.advance_to(end, collect=True):
                if change.emp_account and change.type in trigger_of:
                    accounts.add(change.emp_account)
                    at = change.at.replace(second=0, microsecond=0)
                    triggers.setdefault((at, change.emp_account), trigger_of[change.type])
            fakes = {account: fake_tasks(ctx, account, day, seed, next_id) for account in sorted(accounts)}
            for account, items in fakes.items():
                triggers[(start, account)] = "DAY_START"
                for task in items:
                    if start < task.created <= end:
                        triggers.setdefault((task.created, account), "TASK_NEW")
                    if task.done is not None and start < task.done <= end:
                        triggers.setdefault((task.done, account), "TASK_DONE")
                    if task.reschedule is not None:
                        triggers.setdefault((task.reschedule[0], account), "RESCHEDULE")
            for (at, account), trigger in sorted(triggers.items()):
                if not start <= at <= end:
                    continue
                request = provider.build(WorkloadQuery(at, JobFilter(emp_accounts=(account,))))[0]
                technician = request.technicians[0] if request.technicians else None
                message, source = build_message(
                    ctx, f"M{len(messages) + 1:06d}", at, trigger, account, technician, fakes[account]
                )
                messages.append(message)
                sources.append(source)
    return messages, sources


# ---------------------------------------------------------------- kiểm tra schema

DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
LATLNG = re.compile(r"^(-?\d{1,2}(?:\.\d+)?),(-?\d{1,3}(?:\.\d+)?)$")
AVAILABLE = re.compile(r"^\d{2}:\d{2}-\d{2}:\d{2}(,\d{2}:\d{2}-\d{2}:\d{2})*$")
STAFF_FIELDS = {"staff_id", "staff_account", "staff_location", "latlng", "available", "plots", "current_task"}
TASK_FIELDS = {
    "task_id", "task_group_id", "task_group_name", "task_type_id", "task_type_name", "task_sub_id",
    "task_sub_name", "task_status_id", "task_status_name", "sla", "appointment", "location", "latlng",
    "handle_minutes", "task_plots_id", "staff_plots_id", "staff_role", "block_id",
    "location_id",  # Có trong JSON mẫu, không có trong bảng field.
}
MESSAGE_FIELDS = {"message_id", "planned_at", "trigger", "staff", "tasks"}


def _is_int(value) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def validate(message) -> list[tuple[str, str, str]]:
    """Lỗi so với file API: [(đường dẫn field, lỗi, giá trị)]. Rỗng là hợp lệ."""

    errors: list[tuple[str, str, str]] = []

    def fail(path: str, problem: str, value) -> None:
        errors.append((path, problem, json.dumps(value, ensure_ascii=False)[:120]))

    def need(data: dict, path: str, key: str, check, problem: str, required: bool = True) -> None:
        if key not in data:
            if required:
                fail(f"{path}.{key}", "thiếu field bắt buộc", None)
            return
        if not check(data[key]):
            fail(f"{path}.{key}", problem, data[key])

    def unknown(data: dict, path: str, allowed: set[str]) -> None:
        for key in sorted(set(data) - allowed):
            fail(f"{path}.{key}", "field không có trong file API", data[key])

    def text(value) -> bool:
        return isinstance(value, str)

    def when(value) -> bool:
        if not isinstance(value, str) or not DATETIME.match(value):
            return False
        try:
            datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return False
        return True

    def latlng(value) -> bool:
        match = isinstance(value, str) and LATLNG.match(value)
        return bool(match) and 8 <= float(match[1]) <= 24 and 102 <= float(match[2]) <= 110

    def available(value) -> bool:
        if not isinstance(value, str) or not AVAILABLE.match(value):
            return False
        return all(part[:5] < part[6:] for part in value.split(","))

    if not isinstance(message, dict):
        fail("", "cần object JSON", message)
        return errors
    unknown(message, "", MESSAGE_FIELDS)
    need(message, "", "message_id", lambda v: isinstance(v, str) and v != "", "cần chuỗi khác rỗng")
    need(message, "", "planned_at", when, 'cần "YYYY-MM-DD HH:mm:ss"')
    need(message, "", "trigger", lambda v: v in TRIGGERS, f"cần một trong {', '.join(TRIGGERS)}")

    staff = message.get("staff")
    if not isinstance(staff, dict):
        fail("staff", "cần object", staff)
    else:
        unknown(staff, "staff", STAFF_FIELDS)
        need(staff, "staff", "staff_id", lambda v: isinstance(v, str) and v != "", "cần chuỗi khác rỗng")
        need(staff, "staff", "staff_account", lambda v: isinstance(v, str) and v != "", "cần chuỗi khác rỗng")
        need(staff, "staff", "staff_location", text, "cần chuỗi", required=False)
        need(staff, "staff", "latlng", latlng, 'cần "lat,lng" trong Việt Nam')
        need(staff, "staff", "available", available, 'cần "HH:mm-HH:mm,..." với giờ đầu < giờ cuối')
        plots = staff.get("plots")
        if not isinstance(plots, list) or not plots:
            fail("staff.plots", "cần mảng khác rỗng", plots)
        else:
            for index, plot in enumerate(plots):
                path = f"staff.plots[{index}]"
                if not isinstance(plot, dict):
                    fail(path, "cần object", plot)
                    continue
                unknown(plot, path, {"id", "name", "role", "block_id"})
                need(plot, path, "id", _is_int, "cần số nguyên")
                need(plot, path, "name", text, "cần chuỗi", required=False)
                need(plot, path, "role", lambda v: v in (1, 2), "cần 1 hoặc 2")
                need(plot, path, "block_id", _is_int, "cần số nguyên")
            if sum(plot.get("role") == 1 for plot in plots if isinstance(plot, dict)) != 1:
                fail("staff.plots", "cần đúng một lô chính (role = 1)", [p.get("role") for p in plots if isinstance(p, dict)])
        if "current_task" not in staff:
            fail("staff.current_task", "thiếu field bắt buộc (null nếu không có)", None)
        elif staff["current_task"] is not None:
            current = staff["current_task"]
            if not isinstance(current, dict):
                fail("staff.current_task", "cần object hoặc null", current)
            else:
                unknown(current, "staff.current_task", {"task_id", "task_status_id", "task_type_id"})
                for key in ("task_id", "task_status_id", "task_type_id"):
                    need(current, "staff.current_task", key, _is_int, "cần số nguyên")

    tasks = message.get("tasks")
    if not isinstance(tasks, dict):
        fail("tasks", "cần object 5 khóa", tasks)
        return errors
    if set(tasks) != set(GROUP_KEYS):
        fail("tasks", f"cần đúng 5 khóa {', '.join(GROUP_KEYS)}", sorted(tasks))
    seen: set = set()
    for key in GROUP_KEYS:
        items = tasks.get(key)
        if not isinstance(items, list):
            fail(f"tasks.{key}", "cần mảng ([] nếu không có việc)", items)
            continue
        for index, task in enumerate(items):
            path = f"tasks.{key}[{index}]"
            if not isinstance(task, dict):
                fail(path, "cần object", task)
                continue
            unknown(task, path, TASK_FIELDS)
            for name in ("task_id", "task_group_id", "task_type_id", "task_sub_id", "task_status_id",
                         "task_plots_id", "staff_plots_id", "block_id"):
                need(task, path, name, _is_int, "cần số nguyên", required=name != "task_sub_id")
            for name in ("task_group_name", "task_type_name", "task_status_name"):
                need(task, path, name, text, "cần chuỗi")
            need(task, path, "task_sub_name", text, "cần chuỗi", required=False)
            need(task, path, "appointment", lambda v: v == "" or when(v), 'cần "" hoặc "YYYY-MM-DD HH:mm:ss"', required=False)
            need(task, path, "location", text, "cần chuỗi", required=False)
            # Thiếu tọa độ ("") là hợp lệ: quy tắc cứng 4 loại việc đó khỏi tuyến.
            need(task, path, "latlng", lambda v: v == "" or latlng(v), 'cần "" hoặc "lat,lng" trong Việt Nam')
            need(task, path, "handle_minutes", lambda v: v == "" or (_is_int(v) and v > 0), 'cần "" hoặc số nguyên > 0', required=False)
            need(task, path, "staff_role", lambda v: v in (1, 2), "cần 1 hoặc 2")
            need(task, path, "location_id", _is_int, "cần số nguyên", required=False)
            sla = task.get("sla")
            if not isinstance(sla, dict):
                fail(f"{path}.sla", "cần object {sla_minutes, priority_in_day}", sla)
            else:
                unknown(sla, f"{path}.sla", {"sla_minutes", "priority_in_day"})
                need(sla, f"{path}.sla", "sla_minutes", lambda v: v is None or (_is_int(v) and v > 0), "cần số nguyên > 0 hoặc null")
                need(sla, f"{path}.sla", "priority_in_day", lambda v: v in (1, 2, 3, 4), "cần 1–4")
            if task.get("task_group_name") != key or task.get("task_group_id") != GROUP_IDS[key]:
                fail(path, f"nằm trong nhóm {key} nhưng task_group là", [task.get("task_group_id"), task.get("task_group_name")])
            kind = KIND_BY_NAME.get((key, task.get("task_type_name")))
            if kind is None:
                fail(f"{path}.task_type_name", "không có trong danh mục sheet 05", task.get("task_type_name"))
            else:
                if task.get("task_type_id") != kind.type_id:
                    fail(f"{path}.task_type_id", f"danh mục ghi {kind.type_id}", task.get("task_type_id"))
                if isinstance(sla, dict) and (sla.get("sla_minutes"), sla.get("priority_in_day")) != (kind.sla_minutes, kind.priority):
                    fail(f"{path}.sla", f"danh mục ghi sla_minutes={kind.sla_minutes}, priority_in_day={kind.priority}", sla)
            if task.get("task_id") in seen:
                fail(f"{path}.task_id", "trùng với việc khác trong message", task.get("task_id"))
            seen.add(task.get("task_id"))
    current = staff.get("current_task") if isinstance(staff, dict) else None
    if isinstance(current, dict) and current.get("task_id") in seen:
        fail("staff.current_task.task_id", "việc đang làm vẫn nằm trong tasks", current.get("task_id"))
    return errors


# ---------------------------------------------------------------- tình huống


def _all_tasks(message: dict) -> list[dict]:
    return [task for key in GROUP_KEYS for task in message["tasks"][key]]


def _point(value: str) -> GeoPoint | None:
    if not value:
        return None
    lat, lng = value.split(",")
    return GeoPoint(float(lat), float(lng))


def _window(task: dict) -> tuple[datetime, datetime] | None:
    if not task["appointment"] or task["sla"]["sla_minutes"] is None:
        return None
    start = datetime.strptime(task["appointment"], "%Y-%m-%d %H:%M:%S")
    return start, start + timedelta(minutes=task["sla"]["sla_minutes"])


def _handle(task: dict) -> int:
    if task["handle_minutes"] != "":
        return task["handle_minutes"]
    return KIND_BY_NAME[(task["task_group_name"], task["task_type_name"])].handle[0]


def _cannot_meet_all(message: dict) -> bool:
    """Có hai việc hẹn mà làm theo thứ tự nào cũng trễ một việc (đi 30 km/h, đường chim bay)."""

    now = datetime.strptime(message["planned_at"], "%Y-%m-%d %H:%M:%S")
    booked = [(task, _window(task), _point(task["latlng"])) for task in _all_tasks(message)]
    booked = [item for item in booked if item[1] is not None and item[2] is not None and item[1][1] >= now]

    def fits(first, second) -> bool:
        done = max(first[1][0], now) + timedelta(minutes=_handle(first[0]))
        arrive = done + timedelta(minutes=distance_km(first[2], second[2]) / 30 * 60)
        return arrive <= second[1][1]

    return any(
        not fits(a, b) and not fits(b, a) for i, a in enumerate(booked) for b in booked[i + 1 :]
    )


def _overlap(message: dict) -> bool:
    starts = sorted(window[0] for task in _all_tasks(message) if (window := _window(task)))
    return any(later - earlier < timedelta(minutes=30) for earlier, later in zip(starts, starts[1:]))


def _overdue(message: dict) -> bool:
    now = datetime.strptime(message["planned_at"], "%Y-%m-%d %H:%M:%S")
    return any((window := _window(task)) and window[1] < now for task in _all_tasks(message))


def _outside_plots(message: dict) -> bool:
    own = {plot["id"] for plot in message["staff"]["plots"]}
    return any(task["task_plots_id"] not in own for task in _all_tasks(message))


SCENARIOS = (
    ("Bình thường", "≥ 3 việc, không thiếu tọa độ, không hẹn không kịp, không quá hạn",
     lambda m: len(_all_tasks(m)) >= 3 and not any(t["latlng"] == "" for t in _all_tasks(m))
     and not _cannot_meet_all(m) and not _overdue(m)),
    ("Có việc đang làm", "staff.current_task khác null → khóa đầu tuyến", lambda m: m["staff"]["current_task"] is not None),
    ("Có việc thiếu latlng", "Quy tắc cứng 4: loại khỏi tuyến", lambda m: any(t["latlng"] == "" for t in _all_tasks(m))),
    ("Hẹn chồng giờ", "Hai giờ hẹn cách nhau dưới 30 phút", _overlap),
    ("Không thể kịp mọi hẹn", "Hai việc hẹn, làm thứ tự nào cũng trễ một việc", _cannot_meet_all),
    ("Có việc đã quá hạn khi gọi", "Mốc hẹn cuối < planned_at → ALREADY_BREACHED", _overdue),
    ("KTV có OT", "available có 2 khung giờ", lambda m: "," in m["staff"]["available"]),
    ("Trên 12 việc", "Quá ngưỡng QHĐ chính xác, cần heuristic", lambda m: len(_all_tasks(m)) > 12),
    ("tasks rỗng", "Mã lỗi 422: không có việc để dựng tuyến", lambda m: not _all_tasks(m)),
    ("Đủ 5 nhóm việc", "Cả 5 khóa đều có việc", lambda m: all(m["tasks"][key] for key in GROUP_KEYS)),
    ("Có việc ở lô phụ", "staff_role = 2", lambda m: any(t["staff_role"] == 2 for t in _all_tasks(m))),
    ("Có việc ngoài mọi lô của KTV", "task_plots_id không nằm trong staff.plots", _outside_plots),
    ("Khách đổi giờ", "trigger = RESCHEDULE", lambda m: m["trigger"] == "RESCHEDULE"),
)
SCENARIO_TARGET = 30


# ---------------------------------------------------------------- Excel


def write_review(
    path: Path,
    messages: list[dict],
    sources: list[dict[int, str]],
    ctx: Context,
    days: list[date],
    errors: list[tuple[str, str, str, str]],
    scenarios: list[tuple[str, str, list[str]]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fills = {
        "GIẢ": PatternFill("solid", fgColor="FFF2CC"),
        "SUY": PatternFill("solid", fgColor="DDEBF7"),
        "THẬT": None,
    }
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    book = Workbook()
    book.remove(book.active)

    def sheet(name: str, headers: list[str], widths: list[int], *, start_row: int = 1):
        ws = book.create_sheet(name)
        for column, (header, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(start_row, column, header)
            cell.fill, cell.font, cell.alignment = head_fill, head_font, wrap
            ws.column_dimensions[get_column_letter(column)].width = width
        ws.freeze_panes = ws.cell(start_row + 1, 1)
        return ws

    def finish(ws, start_row: int = 1) -> None:
        if ws.max_row > start_row:
            ws.auto_filter.ref = f"A{start_row}:{get_column_letter(ws.max_column)}{ws.max_row}"

    task_count = sum(len(_all_tasks(message)) for message in messages)
    fake_count = sum(1 for source in sources for kind in source.values() if kind == "FAKE")
    staff_count = len({message["staff"]["staff_account"] for message in messages})
    covered = [name for name, _, ids in scenarios if len(ids) >= SCENARIO_TARGET]
    day_list = ", ".join(day.strftime("%d/%m/%Y") for day in sorted(days))

    # 00 · Tổng quan
    ws = book.create_sheet("00 · Tổng quan")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 95
    section_rows: list[int] = []

    def title(text: str) -> None:
        ws.append([])
        ws.append([text])
        section_rows.append(ws.max_row)

    ws.append(["Dữ liệu giả lập cho AI Gợi ý công việc"])
    ws["A1"].font = title_font
    ws.append([None, "Bản dành cho người review · không cần đọc code"])

    title("1. Vì sao có bộ dữ liệu này")
    ws.append([None, None, "AI Gợi ý công việc nhận danh sách việc của từng KTV rồi sắp xếp thứ tự làm. Hệ thống thật chưa gửi được dữ liệu "
               "theo đúng định dạng API, nên nhóm tạo bộ dữ liệu giả lập, dựa tối đa trên dữ liệu thật, để thử AI trước."])
    ws.append([None, None, "Nhờ anh/chị xem dữ liệu này có giống thực tế không. Nếu giống, nhóm sẽ dùng nó để chạy thử AI "
               "và đo chất lượng sắp xếp trước khi có dữ liệu thật."])

    title("2. Dữ liệu được tạo ra như thế nào")
    steps = [
        ("Bước 1", "Lấy dữ liệu thật",
         "Việc bảo trì tháng 6/2026 của chi nhánh HNI_04: ai phụ trách, địa chỉ khách, lúc tạo phiếu, lúc KTV check-in / checkout "
         "và vị trí check-in (chính là vị trí nhà khách)."),
        ("Bước 2", "Bổ sung 4 nhóm việc còn thiếu",
         "Chưa có dữ liệu của triển khai, thu hồi, hóa đơn, onsite. Mỗi KTV được tạo thêm vài việc mỗi ngày, đặt tại địa chỉ khách thật "
         "trong địa bàn của KTV; SLA và mức ưu tiên theo đúng bảng SLA; thời gian xử lý và giờ hẹn theo các con số ở sheet 01."),
        ("Bước 3", "Xác định địa bàn",
         "Mỗi phường là một địa bàn. Địa bàn chính của KTV là nơi KTV làm nhiều việc nhất; địa bàn hỗ trợ là nơi KTV thường sang giúp."),
        ("Bước 4", "Tái hiện một ngày làm việc",
         "Chạy lại từng ngày theo đúng thứ tự sự kiện thật. Mỗi khi có chuyện xảy ra với một KTV (6h sáng đầu ngày, làm xong một việc, "
         "có việc mới, khách đổi hẹn), chụp lại danh sách việc của KTV lúc đó. Mỗi lần chụp là một bản ghi, đúng bằng một lần gửi cho AI."),
    ]
    for step, name, text in steps:
        ws.append([step, name, text])

    title("3. Quy mô")
    for name, value in [
        ("Chi nhánh", "HNI_04"),
        ("Ngày", f"{day_list} (ngày cao điểm, ngày thường, chủ nhật)"),
        ("Số KTV", f"{staff_count:,}"),
        ("Số lần gửi cho AI", f"{len(messages):,} (sheet 02, mỗi dòng một lần)"),
        ("Số dòng việc", f"{task_count:,} (sheet 04), trong đó {task_count - fake_count:,} việc bảo trì thật và {fake_count:,} việc tạo thêm"),
    ]:
        ws.append([None, name, value])

    title("4. Kết quả tự kiểm tra")
    ws.append([None, "Đúng định dạng file API", "Đạt: mọi bản ghi đều đúng định dạng" if not errors
               else f"Chưa đạt: {len({error[0] for error in errors})} bản ghi sai định dạng"])
    ws.append([None, "Có đủ tình huống để thử AI",
               f"{len(covered)}/{len(scenarios)} tình huống có ít nhất {SCENARIO_TARGET} bản ghi: " + "; ".join(name.lower() for name in covered)])

    title("5. Nhờ anh/chị review")
    for index, text in enumerate([
        "Đọc sheet 01. Mỗi dòng là một điều nhóm phải tự đặt hoặc suy ra; chọn \"Đồng ý\" hoặc \"Sửa\" ở cột Duyệt, ghi con số đúng vào cột Ghi chú.",
        "Lọc vài KTV quen ở sheet 02 và 04, xem số việc mỗi ngày, loại việc, giờ hẹn, địa bàn có giống thực tế không.",
        "Đối chiếu bản ghi mẫu ở sheet 05 với dữ liệu hệ thống đang có (dành cho team hệ thống).",
    ], start=1):
        ws.append([index, None, text])
    ws.append([None, "Tiến độ duyệt", None])
    ws.cell(ws.max_row, 3).value = f'=COUNTIF(\'01 · Giả định\'!G:G,"Đồng ý")&" / {len(ASSUMPTIONS)} dòng đã đồng ý"'

    title("6. Cách đọc")
    for color, name, text in [
        (None, "Ô trắng", "Thật: lấy nguyên từ dữ liệu thật"),
        ("SUY", "Ô xanh", "Suy ra: tính từ dữ liệu thật hoặc từ bảng SLA"),
        ("GIẢ", "Ô vàng", "Tự đặt: con số nhóm tự chọn, cần duyệt ở sheet 01"),
    ]:
        ws.append([None, name, text])
        if color:
            ws.cell(ws.max_row, 2).fill = fills[color]
    for name, text in [
        ("01 · Giả định", "Những điều nhóm tự đặt hoặc suy ra, kèm ô duyệt"),
        ("02 · Input staff", "Mỗi dòng là một lần gửi cho AI: KTV nào, lúc nào, vì sao gửi, vị trí, ca làm, việc đang làm dở"),
        ("03 · Input plots", "Địa bàn chính và địa bàn hỗ trợ của từng KTV, kèm tỷ lệ việc thật trong từng địa bàn"),
        ("04 · Input tasks", "Từng việc trong từng lần gửi, tên cột đúng như file API, ô tô màu theo nguồn"),
        ("05 · JSON mẫu", "Vài bản ghi hoàn chỉnh đúng định dạng API, mỗi bản ghi một tình huống"),
    ]:
        ws.append([None, name, text])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = wrap
    for row in section_rows:
        ws.cell(row, 1).font = Font(bold=True, size=12)

    # 01 · Giả định
    ws = sheet(
        "01 · Giả định",
        ["#", "Chủ đề", "Field trong API", "Cách giả lập", "Ví dụ / con số", "Nguồn", "Duyệt", "Ghi chú của người review"],
        [5, 28, 22, 70, 42, 10, 11, 36],
    )
    source_fill = {"Thật": None, "Suy ra": fills["SUY"], "Tự đặt": fills["GIẢ"]}
    for index, (topic, api_field, rule, example, source) in enumerate(ASSUMPTIONS, start=1):
        ws.append([index, topic, api_field, rule, example, source, "", ""])
        if source_fill[source] is not None:
            ws.cell(ws.max_row, 6).fill = source_fill[source]
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
    choice = DataValidation(type="list", formula1='"Đồng ý,Sửa"', allow_blank=True)
    ws.add_data_validation(choice)
    choice.add(f"G2:G{ws.max_row}")
    finish(ws)

    # 02 · Input staff
    ws = sheet(
        "02 · Input staff",
        ["message_id", "planned_at", "trigger", "staff_id", "staff_account", "staff_location", "latlng",
         "available", "plots (id:role)", "current_task.task_id", "current_task.task_status_id",
         "current_task.task_type_id", "Số việc"],
        [11, 19, 13, 11, 22, 10, 22, 24, 22, 14, 12, 12, 8],
    )
    for message in messages:
        staff, current = message["staff"], message["staff"]["current_task"] or {}
        ws.append([
            message["message_id"], message["planned_at"], message["trigger"], staff["staff_id"],
            staff["staff_account"], staff["staff_location"], staff["latlng"], staff["available"],
            ", ".join(f"{plot['id']}:{plot['role']}" for plot in staff["plots"]),
            current.get("task_id", ""), current.get("task_status_id", ""), current.get("task_type_id", ""),
            len(_all_tasks(message)),
        ])
    finish(ws)

    # 03 · Input plots
    ws = sheet(
        "03 · Input plots",
        ["staff_id", "staff_account", "id", "name", "role", "block_id", "% việc của KTV trong lô (cả tháng)"],
        [11, 22, 8, 22, 6, 9, 16],
    )
    for account in sorted({message["staff"]["staff_account"] for message in messages}):
        for plot, role, share in ctx.staff_plots.get(account, []):
            ws.append([ctx.staff_ids[account][0], account, plot.id, plot.name, role, BLOCK_ID, round(share * 100, 1)])
    finish(ws)

    # 04 · Input tasks
    columns = ["task_id", "task_group_id", "task_group_name", "task_type_id", "task_type_name", "task_sub_id",
               "task_sub_name", "task_status_id", "task_status_name", "sla.sla_minutes", "sla.priority_in_day",
               "appointment", "location", "latlng", "handle_minutes", "task_plots_id", "staff_plots_id",
               "staff_role", "block_id", "location_id"]
    real_source = {
        "task_id": "THẬT", "task_group_id": "SUY", "task_group_name": "SUY", "task_type_id": "GIẢ",
        "task_type_name": "GIẢ", "task_sub_id": "SUY", "task_sub_name": "SUY", "task_status_id": "GIẢ",
        "task_status_name": "GIẢ", "sla.sla_minutes": "SUY", "sla.priority_in_day": "GIẢ", "appointment": "GIẢ",
        "location": "THẬT", "latlng": "THẬT", "handle_minutes": "SUY", "task_plots_id": "SUY",
        "staff_plots_id": "SUY", "staff_role": "SUY", "block_id": "GIẢ", "location_id": "GIẢ",
    }
    fake_source = dict.fromkeys(columns, "GIẢ") | {
        "sla.sla_minutes": "SUY", "sla.priority_in_day": "SUY", "task_plots_id": "SUY",
        "staff_plots_id": "SUY", "staff_role": "SUY",
    }
    ws = sheet(
        "04 · Input tasks",
        ["message_id", "planned_at", "Việc thật / tạo thêm", *columns],
        [11, 19, 10, 12, 7, 12, 7, 18, 7, 9, 7, 10, 8, 8, 19, 40, 22, 9, 9, 9, 7, 8, 9],
    )
    for message, source in zip(messages, sources):
        for task in _all_tasks(message):
            kind = source[task["task_id"]]
            values = {
                key: task["sla"][key[4:]] if key.startswith("sla.") else task[key] for key in columns
            }
            ws.append([message["message_id"], message["planned_at"], "Tạo thêm" if kind == "FAKE" else "Thật",
                       *[("null" if values[key] is None else values[key]) for key in columns]])
            row = ws.max_row
            mapping = fake_source if kind == "FAKE" else real_source
            for column, key in enumerate(columns, start=4):
                label = "SUY" if kind == "REAL_WARD" and key == "latlng" else mapping[key]
                if fills[label] is not None:
                    ws.cell(row, column).fill = fills[label]
    finish(ws)

    # 05 · JSON mẫu
    ws = book.create_sheet("05 · JSON mẫu")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws.append(["Bản ghi mẫu — mỗi tình huống một bản ghi hoàn chỉnh, đúng định dạng file API"])
    ws["A1"].font = title_font
    by_id = {message["message_id"]: message for message in messages}
    shown: set[str] = set()
    for name, _, ids in scenarios:
        pick = next((message_id for message_id in ids if message_id not in shown), None)
        if pick is None:
            continue
        shown.add(pick)
        ws.append([])
        ws.append([f"{name} — {pick}"])
        ws.cell(ws.max_row, 1).font = bold
        for line in json.dumps(by_id[pick], ensure_ascii=False, indent=2).splitlines():
            ws.append([None, line])

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


# ---------------------------------------------------------------- CLI


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m ktv_simulator.fake_worklist",
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--events", type=Path, default=Path("data/sample/events_2026-06_HNI_04.jsonl"))
    parser.add_argument("--checkins", type=Path, default=Path("data/sample/QOS_MAINT_CHECKIN_INFO_utf8.csv"))
    parser.add_argument("--day", action="append", type=date.fromisoformat, dest="days",
                        help="Ngày sinh message (lặp được); mặc định 08/06 cao điểm, 16/06 thường, 21/06 chủ nhật")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--out-dir", type=Path, default=Path("artifacts/fake"))
    args = parser.parse_args(argv)
    days = args.days or [date(2026, 6, 8), date(2026, 6, 16), date(2026, 6, 21)]

    ctx = load_context(args.events, args.checkins)
    messages, sources = generate(args.events, ctx, days, args.seed)
    errors = [(message["message_id"], *error) for message in messages for error in validate(message)]
    scenarios = [
        (name, description, [message["message_id"] for message in messages if check(message)])
        for name, description, check in SCENARIOS
    ]

    args.out_dir.mkdir(parents=True, exist_ok=True)
    jsonl = args.out_dir / "messages.jsonl"
    with jsonl.open("w", encoding="utf-8") as handle:
        for message in messages:
            handle.write(json.dumps(message, ensure_ascii=False) + "\n")
    review = args.out_dir / "FAKE_DATA_REVIEW.xlsx"
    write_review(review, messages, sources, ctx, days, errors, scenarios)

    print(f"{len(messages):,} message, {sum(len(_all_tasks(m)) for m in messages):,} dòng việc, lỗi schema {len(errors)}")
    for name, _, ids in scenarios:
        print(f"  {name:32} {len(ids):6,}{'' if len(ids) >= SCENARIO_TARGET else '  ← chưa đủ'}")
    print(f"Đã ghi {jsonl} và {review}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
