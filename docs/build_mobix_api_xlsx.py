#!/usr/bin/env python3
"""Sinh docs/MOBIX-REPLAN-API.xlsx từ nội dung docs/MOBIX-REPLAN-API-DRAFT.md.

Chạy: .venv/bin/python docs/build_mobix_api_xlsx.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).with_name("MOBIX-REPLAN-API.xlsx")

head_fill = PatternFill("solid", fgColor="1F4E78")
head_font = Font(bold=True, color="FFFFFF")
title_font = Font(bold=True, size=14)
sect_font = Font(bold=True, size=12, color="1F4E78")
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "API Replan"


def row(values, style=None):
    ws.append(values)
    r = ws.max_row
    for c in range(1, len(values) + 1):
        cell = ws.cell(r, c)
        cell.alignment = wrap
        cell.border = border
        if style:
            cell.font = style
    return r


def header(values):
    r = row(values)
    for c in range(1, len(values) + 1):
        ws.cell(r, c).fill = head_fill
        ws.cell(r, c).font = head_font
    return r


def title(text):
    ws.cell(row([text]), 1).font = title_font


def section(text):
    row([text], sect_font)


for col, width in zip("ABCDE", (30, 14, 12, 16, 70)):
    ws.column_dimensions[col].width = width

title("API Mobix yêu cầu Core AI sắp xếp lại (replan) — DRAFT để viết ticket")
row(["Draft, chờ Mobix/OA/Gateway chốt. Không chứa credential. Format tham khảo IVR-API.xlsx nhưng đây là API routing, không phải /onebot/run-task."])
row([])

section("1. Configuration")
header(["Configuration", "Description"])
row(["Method", "GET (có tác dụng phụ: trigger replan)"])
row(["URL", "https://{GATEWAY_URL}/api/v1/staff/{staff_id}/replan"])
row(["Content Type", "Không có request body"])
row(["Authorization", "Bearer <KTV token>; chỉ KTV sở hữu staff_id"])
row(["Response", "JSON acknowledgment — KHÔNG trả route"])
row(["Cache", "Cache-Control: no-store"])
row([])

section("2. Request — Input")
row(["Không có request body. staff_id trên path; tọa độ/metadata trên query/header."])
header(["Tên trường", "Vị trí", "Bắt buộc (Y/N)", "Kiểu dữ liệu", "Mô tả"])
row(["staff_id", "path", "Y", "string", "ID KTV, giữ leading zero. Ví dụ: 00201964"])
row(["latlng", "query", "Y", "string", 'Tọa độ mới Mobix biết, format "lat,lng" WGS84'])
row(["latlng_at", "query", "Y", "string", 'Thời điểm ghi nhận tọa độ, giờ VN "YYYY-MM-DD HH:mm:ss"'])
row(["date", "query", "N", "string", 'Ngày tuyến "YYYY-MM-DD"; bỏ trống = ngày hiện hành'])
row(["X-Request-ID", "header", "N", "string", "Mobix sinh để trace; retry nên dùng lại cùng ID"])
row([])

section("3. Response — Output HTTP")
row(["HTTP chỉ trả acknowledgment; không chứa clusters/schedule/route. Route đi Kafka OUT."])
header(["Tên trường", "Kiểu dữ liệu", "Mô tả"])
row(["success", "bool", "true khi nhận được yêu cầu"])
row(["statuscode", "string", '"202"'])
row(["message", "string", 'Ví dụ "Replan accepted"'])
row(["request_id", "string", "ID yêu cầu (echo X-Request-ID hoặc server sinh)"])
row(["trace_id", "string", "Correlation sang Kafka OUT"])
row([])

header(["HTTP/statuscode", "Ý nghĩa", "Client xử lý"])
row(["202", "Đã nhận replan; kết quả publish Kafka OUT", "Không coi response là route"])
row(["400", "Sai format staff_id/latlng/latlng_at/date", "Không retry nguyên request"])
row(["401", "Token thiếu/hết hạn/sai", "Refresh token rồi gửi lại"])
row(["403", "Token không có quyền với KTV này", "Không retry"])
row(["404", "Chưa có snapshot Kafka IN cho staff/ngày (cần chốt UX)", "Retry sau hoặc báo chưa sẵn sàng"])
row(["429", "Gateway rate limit (nếu bật)", "Chờ theo Retry-After"])
row(["5xx", "Gateway/Kafka/planner lỗi trước khi nhận lệnh", "Retry cùng request ID"])
row([])

section("4. Request example")
row(['curl --get --url "https://{GATEWAY_URL}/api/v1/staff/00201964/replan" \\'])
row(['  --header "Authorization: Bearer ${KTV_TOKEN}" --header "X-Request-ID: 01J..." \\'])
row(['  --data-urlencode "latlng=10.7538419,106.7404162" \\'])
row(['  --data-urlencode "latlng_at=2026-09-10 09:20:00" \\'])
row(['  --data-urlencode "date=2026-09-10"'])
row([])

section("5. Return example (202)")
row(['{ "success": true, "statuscode": "202", "message": "Replan accepted", "request_id": "01J...", "trace_id": "01J..." }'])
row([])

section("6. Kafka / state behavior")
header(["Mục", "Nội dung"])
row(["State cache", "Consumer nền giữ snapshot IN mới nhất mỗi (staff_id, date); GET đọc snapshot này"])
row(["Xử lý GET", "auth → đọc state → overlay latlng/latlng_at → plan() → produce Kafka OUT → 202"])
row(["Thứ tự", "Mobix gọi sau khi OA publish snapshot IN đã cập nhật; MVP không yêu cầu state_version, chấp nhận consumer lag nhỏ"])
row(["Dedup", "X-Request-ID nếu có; nếu không dùng fingerprint (staff + snapshot IN + latlng + latlng_at)"])
row(["OUT cần thêm", "Mỗi TASK trong data.clusters[].schedule[] thêm location (địa chỉ) và latlng (tọa độ) cho Mobix hiển thị/map"])
row([])

section("7. Điểm cần chốt trước khi đóng ticket")
header(["#", "Nội dung"])
row([1, "404 khi chưa có snapshot có đúng UX không, hay muốn 202 + retry_after?"])
row([2, "Tên host STAG/PROD (base URL) và token scope: Mobix lấy KTV token từ đâu?"])
row([3, "Tên field tọa độ thống nhất là latlng (theo API-Goi-y) hay location?"])
row([4, "Policy log/retention: không log latlng thô trong access log"])
row([5, "Mode DEFAULT/SLA/DISTANCE thuộc Phase 8 (reoptimize), không trộn vào API này"])

wb.save(OUT)
print(f"saved {OUT}")
